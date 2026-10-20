import io
from datetime import date
from aiohttp import web
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from database import (
    get_dashboard_overview,
    get_items_by_type,
    get_daily_income_breakdown,
    add_transaction,
    delete_transaction_by_id,
    get_all_monthly_transactions
)

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"

MONTH_NAMES_UZ = {
    1: "Yanvar", 2: "Fevral", 3: "Mart", 4: "Aprel",
    5: "May", 6: "Iyun", 7: "Iyul", 8: "Avgust",
    9: "Sentyabr", 10: "Oktyabr", 11: "Noyabr", 12: "Dekabr"
}

TYPE_NAMES_UZ = {
    "income": "Daromad",
    "expense": "Xarajat",
    "advance": "Avans",
    "yasalgan": "Yasalgan buyum",
    "montaj": "Montaj ishi"
}

async def handle_index(request: web.Request):
    return web.FileResponse(STATIC_DIR / "index.html")

async def handle_overview(request: web.Request):
    user_id = int(request.query.get("user_id", 0))
    today = date.today()
    year = int(request.query.get("year", today.year))
    month = int(request.query.get("month", today.month))
    data = await get_dashboard_overview(user_id, year, month)
    return web.json_response(data)

async def handle_items(request: web.Request):
    user_id = int(request.query.get("user_id", 0))
    trans_type = request.query.get("type", "income")
    today = date.today()
    year = int(request.query.get("year", today.year))
    month = int(request.query.get("month", today.month))
    data = await get_items_by_type(user_id, trans_type, year, month)
    return web.json_response(data)

async def handle_add_entry(request: web.Request):
    try:
        body = await request.json()
        user_id = int(body.get("user_id", 0))
        t_type = body.get("type", "income")
        amount = float(body.get("amount", 0.0))
        category = body.get("category", "")
        description = body.get("description", "")
        custom_date = body.get("date") or date.today().isoformat()

        if amount <= 0:
            return web.json_response({"error": "Noto'g'ri summa"}, status=400)

        tx_id = await add_transaction(user_id, t_type, amount, category, description, custom_date)
        return web.json_response({"success": True, "id": tx_id})
    except Exception as e:
        return web.json_response({"error": str(e)}, status=500)

async def handle_delete_entry(request: web.Request):
    user_id = int(request.query.get("user_id", 0))
    trans_id = int(request.match_info.get("id", 0))
    success = await delete_transaction_by_id(user_id, trans_id)
    return web.json_response({"success": success})

async def handle_export(request: web.Request):
    user_id = int(request.query.get("user_id", 0))
    today = date.today()
    year = int(request.query.get("year", today.year))
    month = int(request.query.get("month", today.month))
    month_name = MONTH_NAMES_UZ.get(month, f"{month}-oy")

    overview = await get_dashboard_overview(user_id, year, month)
    items = await get_all_monthly_transactions(user_id, year, month)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Oylik Xulosa"

    # Styles
    title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    summary_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )

    ws.merge_cells("A1:E1")
    ws["A1"] = f"ISHXONA HISOB-KITOB HISOBOTI: {month_name.upper()} {year}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    summary_rows = [
        ("Jami daromad:", overview["income"]["total"]),
        ("Montaj ishlari jami:", overview["montaj"]["total"]),
        ("Yasalgan buyumlar jami:", overview["yasalgan"]["total"]),
        ("Jami xarajatlar:", overview["expense"]["total"]),
        ("Olingan avanslar:", overview["advance"]["total"]),
        ("Ishxonadan qolgan haq ((Daromad+Montaj) - Avans):", overview["net_salary"]),
        ("Sof foyda ((Daromad+Montaj) - Xarajat):", overview["net_profit"])
    ]

    for idx, (label, val) in enumerate(summary_rows, start=3):
        ws[f"A{idx}"] = label
        ws[f"B{idx}"] = val
        ws[f"A{idx}"].font = bold_font
        ws[f"B{idx}"].font = bold_font
        ws[f"A{idx}"].fill = summary_fill
        ws[f"B{idx}"].fill = summary_fill
        ws[f"B{idx}"].number_format = '#,##0'
        ws[f"A{idx}"].border = thin_border
        ws[f"B{idx}"].border = thin_border

    # Table of all items
    start_row = 12
    headers = ["№", "Sana", "Bo'lim / Turi", "Summa (so'm)", "Kategoriya / Nomi", "Izoh / Tafsilot"]
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    current_row = start_row + 1
    for idx, item in enumerate(items, start=1):
        t_type = item['type']
        t_type_name = TYPE_NAMES_UZ.get(t_type, t_type.capitalize())
        row_values = [
            idx,
            item['date'],
            t_type_name,
            item['amount'],
            item['category'] or "-",
            item['description'] or "-"
        ]
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx in [1, 2, 3]:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx == 4:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")
        current_row += 1

    col_widths = {1: 6, 2: 14, 3: 18, 4: 18, 5: 24, 6: 35}
    for col_idx, width in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Hisobot_{month_name}_{year}.xlsx"
    return web.Response(
        body=output.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

async def handle_daily_income(request: web.Request):
    user_id = int(request.query.get("user_id", 0))
    today = date.today()
    year = int(request.query.get("year", today.year))
    month = int(request.query.get("month", today.month))
    data = await get_daily_income_breakdown(user_id, year, month)
    return web.json_response(data)

def create_web_app() -> web.Application:
    app = web.Application()
    app.router.add_get("/", handle_index)
    app.router.add_static("/static", path=STATIC_DIR, name="static")
    app.router.add_get("/api/overview", handle_overview)
    app.router.add_get("/api/items", handle_items)
    app.router.add_get("/api/daily-income", handle_daily_income)
    app.router.add_post("/api/entry", handle_add_entry)
    app.router.add_delete("/api/entry/{id}", handle_delete_entry)
    app.router.add_get("/api/export", handle_export)
    return app
